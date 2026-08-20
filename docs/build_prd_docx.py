from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "熊崽电竞-和平精英C2C游戏陪玩交易平台-PRD-V2.0.md"
OUTPUT = ROOT / "docs" / "熊崽电竞-和平精英C2C游戏陪玩交易平台-PRD-V2.0.docx"
LOGO = ROOT / "public" / "assets" / "xiongzai-logo.png"


# compact_reference_guide preset resolved tokens.
PAGE_WIDTH_IN = 8.5
PAGE_HEIGHT_IN = 11
MARGIN_IN = 1.0
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
BODY_FONT_ASCII = "Arial Unicode MS"
BODY_FONT_CJK = "Arial Unicode MS"
BODY_SIZE = 11
BODY_AFTER = 6
BODY_LINE = 1.25
H1 = {"size": 16, "color": "2E74B5", "before": 18, "after": 10}
H2 = {"size": 13, "color": "2E74B5", "before": 14, "after": 7}
H3 = {"size": 12, "color": "1F4D78", "before": 10, "after": 5}
BLUE = "155E9F"  # named brand override for cover metadata only
LIGHT_BLUE = "E8EEF5"
LIGHTER_BLUE = "F4F8FC"
ORANGE = "FF5A24"  # named brand override for cover kicker only
TEXT = "181A1F"
MUTED = "68717D"
BORDER = "D7E0E8"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size=None, bold=None, color=None, italic=None, mono=False) -> None:
    ascii_font = "Menlo" if mono else BODY_FONT_ASCII
    cjk_font = BODY_FONT_CJK
    run.font.name = ascii_font
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), ascii_font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), ascii_font)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cjk_font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_spacing(paragraph, before=0, after=BODY_AFTER, line=BODY_LINE) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_inline_markdown(paragraph, text: str, size=BODY_SIZE, color=TEXT) -> None:
    # Handles bold and inline code without relying on Word defaults.
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=max(9, size - 1), color="1F4D78", mono=True)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "EEF3F7")
            run._element.get_or_add_rPr().append(shd)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, color=color)


def paragraph_border_bottom(paragraph, color=BORDER, size=8, space=7) -> None:
    p = paragraph._p
    p_pr = p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def set_table_width(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def table_widths(rows: list[list[str]]) -> list[int]:
    cols = len(rows[0])
    if cols == 2:
        first = max(len(r[0]) for r in rows)
        second = max(len(r[1]) for r in rows)
        if first <= 10 and second > first * 2:
            return [2160, 7200]
        return [3600, 5760]
    if cols == 3:
        return [2200, 3260, 3900]
    if cols == 4:
        return [1900, 1900, 2780, 2780]
    widths = [CONTENT_WIDTH_DXA // cols] * cols
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_row_together(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_table(doc, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = len(rows[0])
    table = doc.add_table(rows=0, cols=cols)
    table.style = "Table Grid"
    for row_idx, values in enumerate(rows):
        row = table.add_row()
        keep_row_together(row)
        for col_idx, value in enumerate(values):
            cell = row.cells[col_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            set_paragraph_spacing(p, before=0, after=0, line=1.15)
            add_inline_markdown(p, value, size=9.2 if cols >= 4 else 9.6)
            if row_idx == 0:
                set_cell_shading(cell, LIGHT_BLUE)
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string("1F4D78")
        if row_idx == 0:
            set_repeat_table_header(row)
    set_table_width(table, table_widths(rows))
    table.rows[0].height = None
    after = doc.add_paragraph()
    set_paragraph_spacing(after, before=0, after=5, line=1.0)


def add_callout(doc, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHTER_BLUE)
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, before=0, after=0, line=1.2)
    add_inline_markdown(p, text, size=10.2, color="1F4D78")
    set_table_width(table, [CONTENT_WIDTH_DXA])
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def add_code_block(doc, lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F2F4F7")
    p = cell.paragraphs[0]
    set_paragraph_spacing(p, before=0, after=0, line=1.08)
    for idx, line in enumerate(lines):
        run = p.add_run(line)
        set_run_font(run, size=8.6, color="29323D", mono=True)
        if idx < len(lines) - 1:
            run.add_break()
    set_table_width(table, [CONTENT_WIDTH_DXA])
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT_ASCII
    normal._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT_ASCII)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT_ASCII)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT_CJK)
    normal.font.size = Pt(BODY_SIZE)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    set_paragraph_spacing(normal.paragraph_format._parent if False else doc.add_paragraph(), 0, 0, 1.0)
    # Remove the helper paragraph created above.
    helper = doc.paragraphs[-1]
    helper._element.getparent().remove(helper._element)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(BODY_AFTER)
    normal.paragraph_format.line_spacing = BODY_LINE

    for style_name, token in (("Heading 1", H1), ("Heading 2", H2), ("Heading 3", H3)):
        style = doc.styles[style_name]
        style.font.name = BODY_FONT_ASCII
        style._element.rPr.rFonts.set(qn("w:ascii"), BODY_FONT_ASCII)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), BODY_FONT_ASCII)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT_CJK)
        style.font.size = Pt(token["size"])
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(token["color"])
        style.paragraph_format.space_before = Pt(token["before"])
        style.paragraph_format.space_after = Pt(token["after"])
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True



def create_numbering(doc: Document, kind: str, start: int = 1) -> int:
    """Create a real single-level OOXML numbering definition."""
    root = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in root.findall(qn("w:abstractNum"))
        if node.get(qn("w:abstractNumId")) is not None
    ]
    num_ids = [
        int(node.get(qn("w:numId")))
        for node in root.findall(qn("w:num"))
        if node.get(qn("w:numId")) is not None
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start_node = OxmlElement("w:start")
    start_node.set(qn("w:val"), str(start))
    lvl.append(start_node)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    lvl.append(lvl_text)
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    lvl.append(lvl_jc)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    p_pr.append(ind)
    lvl.append(p_pr)
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), BODY_FONT_ASCII)
    fonts.set(qn("w:hAnsi"), BODY_FONT_ASCII)
    fonts.set(qn("w:eastAsia"), BODY_FONT_CJK)
    r_pr.append(fonts)
    lvl.append(r_pr)
    abstract.append(lvl)
    root.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    root.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = BODY_LINE


def add_list_item(doc: Document, text: str, num_id: int, size=BODY_SIZE):
    p = doc.add_paragraph()
    apply_numbering(p, num_id)
    add_inline_markdown(p, text, size=size)
    return p


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    fld_char = OxmlElement("w:fldChar")
    fld_char.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char, instr, separate, text, end])


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(PAGE_WIDTH_IN)
    section.page_height = Inches(PAGE_HEIGHT_IN)
    section.top_margin = Inches(MARGIN_IN)
    section.bottom_margin = Inches(MARGIN_IN)
    section.left_margin = Inches(MARGIN_IN)
    section.right_margin = Inches(MARGIN_IN)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_spacing(p, 0, 0, 1.0)
    run = p.add_run("熊崽电竞｜产品需求文档 PRD V2.0")
    set_run_font(run, size=8.5, color=MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_paragraph_spacing(fp, 0, 0, 1.0)
    run = fp.add_run("内部开发基线  ·  ")
    set_run_font(run, size=8.5, color=MUTED)
    add_field(fp, "PAGE")


def add_cover(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=8, after=14, line=1.0)
    if LOGO.exists():
        p.add_run().add_picture(str(LOGO), width=Inches(2.35))

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(kicker, 0, 6, 1.0)
    run = kicker.add_run("PRODUCT REQUIREMENTS DOCUMENT")
    set_run_font(run, size=9.5, bold=True, color=ORANGE)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(title, 0, 5, 1.05)
    run = title.add_run("熊崽电竞")
    set_run_font(run, size=29, bold=True, color=BLUE)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(subtitle, 0, 18, 1.15)
    run = subtitle.add_run("和平精英 C2C 游戏陪玩交易平台\n第一阶段 MVP 产品需求文档")
    set_run_font(run, size=17, bold=True, color=TEXT)

    meta_rows = [
        ["文档版本", "V2.0"],
        ["产品形态", "买家 App + 同账号陪玩中心 + 管理员 Web 后台"],
        ["游戏范围", "仅《和平精英》"],
        ["UI 基线", "熊崽电竞浅色、年轻潮流、轻电竞视觉体系"],
        ["文档状态", "开发基线候选版"],
        ["编制日期", "2026-08-18"],
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for label, value in meta_rows:
        row = table.add_row()
        for idx, cell_text in enumerate((label, value)):
            cell = row.cells[idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if idx == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            p = cell.paragraphs[0]
            set_paragraph_spacing(p, 0, 0, 1.1)
            run = p.add_run(cell_text)
            set_run_font(run, size=9.5, bold=idx == 0, color="1F4D78" if idx == 0 else TEXT)
            set_cell_margins(cell, top=100, bottom=100)
    set_table_width(table, [1800, 7560])

    rule = doc.add_paragraph()
    set_paragraph_spacing(rule, 12, 10, 1.0)
    paragraph_border_bottom(rule, color=BLUE, size=18, space=3)

    lead = doc.add_paragraph()
    lead.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(lead, 0, 0, 1.25)
    run = lead.add_run("让个人陪玩能够被看见、被搜索、被咨询、被购买，\n并通过稳定履约获得长期客户。")
    set_run_font(run, size=11.5, bold=True, color="1F4D78")

    doc.add_page_break()


def add_contents(doc: Document, h1s: list[str]) -> None:
    h = doc.add_paragraph(style="Heading 1")
    h.add_run("目录与阅读导航")
    p = doc.add_paragraph()
    set_paragraph_spacing(p, 0, 9, BODY_LINE)
    add_inline_markdown(p, "本文按产品、页面、状态机、后端、数据库、文件存储、安全和验收组织。以下为一级章节：")
    toc_num_id = create_numbering(doc, "decimal", start=1)
    for title in h1s:
        clean_title = re.sub(r"^\d+\.\s*", "", title)
        p = add_list_item(doc, clean_title, toc_num_id, size=9.5)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.05
    doc.add_page_break()


def parse_markdown(doc: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith("# 1."))
    lines = lines[start:]

    i = 0
    first_h1 = True
    bullet_num_id = create_numbering(doc, "bullet", start=1)
    current_decimal_num_id = None
    previous_was_decimal = False
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped or stripped == "---":
            previous_was_decimal = False
            current_decimal_num_id = None
            i += 1
            continue

        if stripped.startswith("```"):
            previous_was_decimal = False
            current_decimal_num_id = None
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i].rstrip("\n"))
                i += 1
            add_code_block(doc, block)
            i += 1
            continue

        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|?\s*:?-+", lines[i + 1]):
            previous_was_decimal = False
            current_decimal_num_id = None
            table_lines = [stripped]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = []
            for table_line in table_lines:
                rows.append([cell.strip() for cell in table_line.strip("|").split("|")])
            add_table(doc, rows)
            continue

        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            previous_was_decimal = False
            current_decimal_num_id = None
            level = len(heading.group(1))
            title = heading.group(2)
            if level == 1:
                if not first_h1:
                    doc.add_page_break()
                first_h1 = False
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline_markdown(p, title, size={1: H1["size"], 2: H2["size"], 3: H3["size"]}[level], color={1: H1["color"], 2: H2["color"], 3: H3["color"]}[level])
            for run in p.runs:
                run.bold = True
            i += 1
            continue

        if stripped.startswith(">"):
            previous_was_decimal = False
            current_decimal_num_id = None
            add_callout(doc, stripped.lstrip("> "))
            i += 1
            continue

        if re.match(r"^-\s+", stripped):
            previous_was_decimal = False
            current_decimal_num_id = None
            add_list_item(doc, re.sub(r"^-\s+", "", stripped), bullet_num_id)
            i += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            source_number = int(re.match(r"^(\d+)\.", stripped).group(1))
            if not previous_was_decimal or current_decimal_num_id is None:
                current_decimal_num_id = create_numbering(doc, "decimal", start=source_number)
            add_list_item(doc, re.sub(r"^\d+\.\s+", "", stripped), current_decimal_num_id)
            previous_was_decimal = True
            i += 1
            continue

        previous_was_decimal = False
        current_decimal_num_id = None
        p = doc.add_paragraph()
        set_paragraph_spacing(p)
        add_inline_markdown(p, stripped)
        i += 1


def set_core_properties(doc: Document) -> None:
    props = doc.core_properties
    props.title = "熊崽电竞｜和平精英 C2C 游戏陪玩交易平台 PRD V2.0"
    props.subject = "第一阶段 MVP 产品需求文档"
    props.author = "熊崽电竞产品团队"
    props.keywords = "熊崽电竞, 和平精英, C2C, 游戏陪玩, PRD, MVP"
    props.comments = "以现有 UI 与原 PRD V1.0 为基础重构的跨端开发基线"


def main() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    h1s = [m.group(1) for line in markdown.splitlines() if (m := re.match(r"^#\s+(.+)$", line)) and re.match(r"\d+\.", m.group(1))]

    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    set_core_properties(doc)
    add_cover(doc)
    add_contents(doc, h1s)
    parse_markdown(doc, markdown)

    # Widow/orphan control and consistent paragraph defaults.
    for paragraph in doc.paragraphs:
        p_pr = paragraph._p.get_or_add_pPr()
        widow = p_pr.find(qn("w:widowControl"))
        if widow is None:
            widow = OxmlElement("w:widowControl")
            p_pr.append(widow)
        widow.set(qn("w:val"), "true")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
